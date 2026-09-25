"""Render fenced code as editable Word runs with Pygments token colors."""

from __future__ import annotations

from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from pygments import lex
from pygments.lexers import get_lexer_by_name
from pygments.styles import get_style_by_name
from pygments.util import ClassNotFound

from markdown_docx.models import CodeBlock

CODE_STYLE = get_style_by_name("default")


def render_code_block(paragraph: Paragraph, block: CodeBlock) -> None:
    """Add code to a paragraph without changing its source text or editability."""
    code = block.text.removesuffix("\n")
    if not code:
        return
    if not block.language:
        paragraph.add_run(code)
        return
    try:
        lexer = get_lexer_by_name(block.language, stripnl=False, ensurenl=False, tabsize=0)
    except ClassNotFound:
        paragraph.add_run(code)
        return
    tokens = list(lex(code, lexer))
    if "".join(value for _, value in tokens) != code:
        paragraph.add_run(code)
        return
    for kind, value in tokens:
        if not value:
            continue
        run = paragraph.add_run(value)
        if value.isspace():
            continue
        style = CODE_STYLE.style_for_token(kind)
        color = style["color"]
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if style["bold"]:
            run.bold = True
        if style["italic"]:
            run.italic = True
        if style["underline"]:
            run.underline = True
