from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from markdown_docx.errors import UnsupportedFeatureError
from markdown_docx.models import ListContentBlock, ListParagraphBlock, ParagraphBlock, ThematicBreakBlock
from markdown_docx.parser import parse_document
from markdown_docx.renderer import render_docx


def parse(source: str):
    return parse_document(source, input_path=None, source_name="input.md")


def render(source: str, tmp_path: Path):
    output = tmp_path / "out.docx"
    render_docx(parse(source), output, template_path=None, base_dir=tmp_path, allow_remote_images=False)
    return Document(output)


def test_thematic_breaks_use_word_paragraph_bottom_borders(tmp_path: Path) -> None:
    document = render("Before\n\n***\n\nAfter\n", tmp_path)
    assert [paragraph.text for paragraph in document.paragraphs] == ["Before", "", "After"]
    rule = document.paragraphs[1]
    assert len(rule._p.xpath("./w:pPr/w:pBdr/w:bottom")) == 1
    assert rule._p.xpath("./w:pPr/w:pBdr/w:bottom/@w:val") == ["single"]


def test_thematic_breaks_work_inside_quotes_and_list_items(tmp_path: Path) -> None:
    source = "- Item\n\n  ---\n\n> ***\n"
    blocks = parse(source).blocks
    assert isinstance(blocks[1], ListContentBlock)
    assert isinstance(blocks[1].content, ThematicBreakBlock)
    assert isinstance(blocks[2], ThematicBreakBlock)
    assert blocks[2].quote_depth == 1
    document = render(source, tmp_path)
    assert len(document.paragraphs[1]._p.xpath("./w:pPr/w:pBdr/w:bottom")) == 1
    assert len(document.paragraphs[2]._p.xpath("./w:pPr/w:pBdr/w:bottom")) == 1
    assert document.paragraphs[1].paragraph_format.left_indent is not None
    assert document.paragraphs[2].paragraph_format.left_indent is not None


@pytest.mark.parametrize("source", ["- one\n-\n- two\n", "1. one\n2.\n3. three\n"])
def test_empty_list_items_keep_native_numbering(source: str, tmp_path: Path) -> None:
    items = [block for block in parse(source).blocks if isinstance(block, ListParagraphBlock)]
    assert len(items) == 3
    assert items[1].fragments == []
    document = render(source, tmp_path)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "one",
        "",
        "two" if source.startswith("-") else "three",
    ]
    ids = [paragraph._p.pPr.numPr.numId.val for paragraph in document.paragraphs]
    assert ids[0] == ids[1] == ids[2]


def test_empty_link_destination_keeps_formatted_label_without_hyperlink(tmp_path: Path) -> None:
    document = render("Before [**label**]() and [more][empty].\n\n[empty]: <>\n", tmp_path)
    paragraph = document.paragraphs[0]
    assert paragraph.text == "Before label and more."
    assert paragraph.hyperlinks == []
    assert any(run.bold and run.text == "label" for run in paragraph.runs)


def test_ordinary_html_comments_are_ignored_but_html_tags_are_rejected(tmp_path: Path) -> None:
    document = render("Before<!-- inline -->after\n\n<!-- block -->\n\n> <!-- quote -->\n", tmp_path)
    assert [paragraph.text for paragraph in document.paragraphs] == ["Beforeafter", ""]
    assert document.paragraphs[1].style.name == "Quote"
    with pytest.raises(UnsupportedFeatureError):
        parse("<em>visible HTML</em>\n")


def test_empty_blockquote_creates_a_blank_quote_paragraph(tmp_path: Path) -> None:
    blocks = parse(">\n").blocks
    assert len(blocks) == 1
    assert isinstance(blocks[0], ParagraphBlock)
    assert blocks[0].quote_depth == 1
    assert blocks[0].fragments == []
    document = render(">\n", tmp_path)
    assert len(document.paragraphs) == 1
    assert document.paragraphs[0].style.name == "Quote"
    assert document.paragraphs[0].text == ""
