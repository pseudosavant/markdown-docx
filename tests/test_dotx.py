from __future__ import annotations

import io
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT

from markdown_docx.cli import main
from markdown_docx.errors import TemplateError
from markdown_docx.parser import parse_document
from markdown_docx.renderer import render_docx
from markdown_docx.template import inspect_template, load_template

FIXTURES = Path(__file__).parent / "fixtures"


def invoke(args: list[str]) -> tuple[int, str, str]:
    stdout = io.StringIO()
    stderr = io.StringIO()
    code = main(args, stdin=io.StringIO(), stdout=stdout, stderr=stderr)
    return code, stdout.getvalue(), stderr.getvalue()


def test_blank_dotx_renders_valid_docx_and_preserves_source(tmp_path: Path) -> None:
    template = FIXTURES / "blank-template-word.dotx"
    original = template.read_bytes()
    model = parse_document("# Report\n\n3. Third\n4. Fourth\n", input_path=None, source_name="input.md")
    output = tmp_path / "report.docx"
    render_docx(model, output, template_path=template, base_dir=tmp_path, allow_remote_images=False)
    document = Document(output)
    assert document.part.content_type == CT.WML_DOCUMENT_MAIN
    assert document.paragraphs[0].text == "Report"
    assert template.read_bytes() == original
    with ZipFile(template) as source, ZipFile(output) as rendered:
        assert rendered.read("word/theme/theme1.xml") == source.read("word/theme/theme1.xml")
        assert CT.WML_DOCUMENT_MAIN.encode() in rendered.read("[Content_Types].xml")
        assert CT.WML_TEMPLATE_MAIN.encode() not in rendered.read("[Content_Types].xml")


@pytest.mark.parametrize("mode", ["--inspect-template", "--list-styles", "--list-table-styles"])
def test_dotx_inspection_and_style_listing(mode: str) -> None:
    template = FIXTURES / "blank-template-word.dotx"
    code, stdout, stderr = invoke([mode, "--template", str(template), "--json"])
    assert code == 0
    assert json.loads(stdout)["ok"] is True
    assert stderr == ""


def test_populated_dotx_is_still_rejected() -> None:
    template = FIXTURES / "rich-template-word.dotx"
    details = inspect_template(template)
    assert details["valid"] is False
    assert "body contains text" in details["errors"]
    assert "section 1 header is not empty" in details["errors"]
    assert "section 1 footer is not empty" in details["errors"]
    with pytest.raises(TemplateError) as error:
        load_template(template)
    assert error.value.context.code == "template_not_blank"


@pytest.mark.parametrize("suffix", [".docm", ".dotm"])
def test_macro_enabled_suffixes_are_rejected_consistently(tmp_path: Path, suffix: str) -> None:
    template = tmp_path / f"template{suffix}"
    template.write_bytes((FIXTURES / "blank-template-word.dotx").read_bytes())
    for operation in (load_template, inspect_template):
        with pytest.raises(TemplateError) as error:
            operation(template)
        assert error.value.context.code == "unsupported_feature"


def test_dotx_cli_render(tmp_path: Path) -> None:
    source = tmp_path / "report.md"
    source.write_text("# Report\n\nEditable body text.", encoding="utf-8")
    code, stdout, stderr = invoke([str(source), "--template", str(FIXTURES / "blank-template-word.dotx"), "--json"])
    assert code == 0
    assert json.loads(stdout)["ok"] is True
    assert stderr == ""
    assert Document(source.with_suffix(".docx")).paragraphs[-1].text == "Editable body text."
