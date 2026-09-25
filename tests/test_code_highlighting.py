from __future__ import annotations

from pathlib import Path

from docx import Document

from markdown_docx.models import CodeBlock, ListContentBlock
from markdown_docx.parser import parse_document
from markdown_docx.renderer import render_docx


def parse(source: str):
    return parse_document(source, input_path=None, source_name="input.md")


def render(source: str, tmp_path: Path):
    output = tmp_path / "out.docx"
    render_docx(parse(source), output, template_path=None, base_dir=tmp_path, allow_remote_images=False)
    return Document(output)


def test_fence_language_is_retained_in_all_block_contexts() -> None:
    blocks = parse(
        "```python title=example.py\nprint(1)\n```\n\n> ```css\n> p { color: red }\n> ```\n\n- item\n\n  ```py\n  x = 1\n  ```\n"
    ).blocks
    assert isinstance(blocks[0], CodeBlock)
    assert blocks[0].language == "python"
    assert isinstance(blocks[1], CodeBlock)
    assert blocks[1].language == "css"
    assert isinstance(blocks[3], ListContentBlock)
    assert isinstance(blocks[3].content, CodeBlock)
    assert blocks[3].content.language == "py"


def test_python_highlighting_creates_colored_editable_runs(tmp_path: Path) -> None:
    document = render('```python\ndef greet(name):\n    print("hi", name)\n```\n', tmp_path)
    paragraph = document.paragraphs[0]
    assert paragraph.text == 'def greet(name):\n    print("hi", name)'
    assert paragraph.style.name == "Code Block"
    assert any(run.text == "def" and run.font.color.rgb is not None for run in paragraph.runs)
    assert any(run.text == "hi" and run.font.color.rgb is not None for run in paragraph.runs)
    assert all(run.font.color.rgb is None for run in paragraph.runs if run.text.isspace())


def test_css_and_alias_highlighting_work_in_quotes_and_lists(tmp_path: Path) -> None:
    source = "> ```css\n> p { color: red }\n> ```\n\n- item\n\n  ```py\n  def f():\n      return 1\n  ```\n"
    document = render(source, tmp_path)
    css, _, python = document.paragraphs
    assert css.text == "p { color: red }"
    assert python.text == "def f():\n    return 1"
    assert any(run.font.color.rgb is not None for run in css.runs)
    assert any(run.font.color.rgb is not None for run in python.runs)
    assert css.paragraph_format.left_indent is not None
    assert python.paragraph_format.left_indent is not None


def test_unknown_and_missing_languages_keep_plain_code(tmp_path: Path) -> None:
    source = "```made-up-language\nprint('plain')\n```\n\n```\nplain\n```\n\n    indented code\n"
    blocks = parse(source).blocks
    assert [block.language for block in blocks if isinstance(block, CodeBlock)] == ["made-up-language", None, None]
    document = render(source, tmp_path)
    assert [paragraph.text for paragraph in document.paragraphs] == ["print('plain')", "plain", "indented code"]
    assert all(len(paragraph.runs) == 1 for paragraph in document.paragraphs)
    assert all(paragraph.runs[0].font.color.rgb is None for paragraph in document.paragraphs)


def test_highlighting_preserves_blank_lines_and_tabs(tmp_path: Path) -> None:
    document = render("```python\nx = 1\n\n\tprint(x)\n\n```\n", tmp_path)
    assert document.paragraphs[0].text == "x = 1\n\n\tprint(x)\n"
