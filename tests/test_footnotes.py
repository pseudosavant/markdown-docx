from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from markdown_docx.errors import MarkdownDocxError
from markdown_docx.parser import parse_document
from markdown_docx.renderer import render_docx


def parse(source: str):
    return parse_document(source, input_path=None, source_name="notes.md")


@pytest.mark.parametrize(
    "source,code,line",
    [
        ("Intro\n\nMissing[^x].", "footnote_undefined", 3),
        ("Intro\nMissing[^x].", "footnote_undefined", 2),
        ("One[^x].\n\n[^x]: First\n\n[^x]: Duplicate", "footnote_duplicate", 5),
        ("One[^x].\n\nAgain[^x].\n\n[^x]: Note", "footnote_repeated_reference", 3),
        ("[^x]: Unused", "footnote_unused", 1),
        ("One[^x].\n\n[^x]: - Nested list", "footnote_content_unsupported", 3),
        ("One[^x].\n\n[^x]: Nested[^y]\n\n[^y]: Note", "footnote_content_unsupported", 3),
        ("One[^x].\n\n[^x]: ![Image](image.png)", "footnote_content_unsupported", 3),
        ("![One[^x]](image.png)\n\n[^x]: Note", "footnote_reference_unsupported", 1),
    ],
)
def test_footnote_diagnostics(source: str, code: str, line: int) -> None:
    with pytest.raises(MarkdownDocxError) as error:
        parse(source)
    assert error.value.context.code == code
    assert error.value.context.line == line
    assert error.value.context.input_path == "notes.md"


def test_code_and_escaped_markers_are_plain_text() -> None:
    model = parse("`[^code]` and \\[^escaped].\n\n```\n[^definition]: code\n```")
    assert not model.footnotes


def test_forward_references_multiline_notes_and_exact_labels() -> None:
    model = parse("One[^Café] and two[^café].\n\n[^Café]: First **bold**\n\n    Second *italic*.\n\n[^café]: Other")
    assert list(model.footnotes) == ["Café", "café"]
    first = model.footnotes["Café"]
    assert len(first.paragraphs) == 2
    assert first.paragraphs[0].fragments[-1].bold
    assert first.paragraphs[1].fragments[1].italic


def test_empty_note_is_supported() -> None:
    model = parse("Text[^empty].\n\n[^empty]:")
    assert model.footnotes["empty"].paragraphs[0].fragments == []


def test_native_footnotes_render_with_links_and_all_body_contexts(tmp_path: Path) -> None:
    model = parse(
        "# Heading[^heading]\n\nText[^text] after.\n\n> Quote[^quote]\n\n"
        "1. Item[^list]\n\n| Header |\n| --- |\n| Cell[^cell] |\n\n"
        "[^heading]: Heading note\n\n[^text]: **Bold** and [website](https://example.com).\n\n"
        "    *Second* paragraph with [heading](#heading).\n\n"
        "[^quote]: Quote note\n\n[^list]: List note\n\n[^cell]: Cell note\n"
    )
    target = tmp_path / "notes.docx"
    render_docx(model, target, template_path=None, base_dir=tmp_path, allow_remote_images=False)
    document = Document(target)
    assert len(document.footnotes) == 5
    assert document.paragraphs[1].text == "Text after."
    assert document.paragraphs[1].runs[-1].text == " after."
    note = document.footnotes.get(2)
    assert note is not None
    assert note.paragraphs[0].runs[3].bold
    assert note.paragraphs[0].hyperlinks[0].address == "https://example.com"
    assert note.paragraphs[1].hyperlinks[0].fragment
    assert document.tables[0].cell(1, 0).paragraphs[0].runs[-1].footnote_ids == (5,)
    part = document.part.part_related_by(RT.FOOTNOTES)
    assert any(rel.reltype == RT.HYPERLINK for rel in part.rels.values())
    assert not any(rel.reltype == RT.HYPERLINK for rel in document.part.rels.values())
    with ZipFile(target) as package:
        xml = package.read(str(part.partname).lstrip("/"))
        assert b"<w:footnoteRef" in xml
        assert b"<w:fldChar" not in xml
