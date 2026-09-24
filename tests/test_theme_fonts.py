from __future__ import annotations

from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.part import Part
from docx.shared import Pt

from markdown_docx.assets import default_template_bytes
from markdown_docx.errors import TemplateError
from markdown_docx.parser import parse_document
from markdown_docx.renderer import render_docx
from markdown_docx.template import load_template

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def w(name: str) -> str:
    return f"{{{NS['w']}}}{name}"


def canonical(element: ET.Element) -> str:
    return ET.canonicalize(ET.tostring(element, encoding="unicode").strip(), strip_text=True)


def render_fonts(tmp_path: Path, fonts: str, *, template: Path | None = None, styles: str = "") -> Path:
    metadata = f"<!-- markdown-docx\ndocument:\n  fonts:\n{fonts}{styles}-->\n\n" if fonts else ""
    headings = "\n\n".join(f"{'#' * level} Heading {level} **bold** *italic* café" for level in range(1, 7))
    source = (
        metadata
        + headings
        + "\n\n"
        + (
            "Body **bold** *italic* café [link](https://example.com) and `inline code`.\n\n"
            "> Quote café\n\n"
            "- Bullet\n  - Nested bullet\n    - Deep bullet\n\n"
            "1. Number\n   1. Nested number\n      1. Deep number\n\n"
            "| Header |\n| --- |\n| Cell café |\n\n```text\nBlock code\n```\n"
        )
    )
    model = parse_document(source, input_path=tmp_path / "fonts.md", source_name="fonts.md")
    output = tmp_path / "fonts.docx"
    render_docx(model, output, template_path=template, base_dir=tmp_path, allow_remote_images=False)
    return output


def xml(path: Path, part: str) -> ET.Element:
    with ZipFile(path) as package:
        return ET.fromstring(package.read(part))


def theme_names(theme: ET.Element) -> tuple[str, str]:
    return tuple(
        theme.find(f"a:themeElements/a:fontScheme/a:{role}Font/a:latin", NS).get("typeface")
        for role in ("major", "minor")
    )


def effective_font(run: ET.Element, paragraph: ET.Element, styles: ET.Element, theme: ET.Element, slot: str) -> str:
    """Resolve the saved Latin font cascade independently of the renderer."""
    by_id = {style.get(w("styleId")): style for style in styles.findall("w:style", NS)}
    candidates = [run.find("w:rPr/w:rFonts", NS)]
    for index, reference in enumerate((run.find("w:rPr/w:rStyle", NS), paragraph.find("w:pPr/w:pStyle", NS))):
        if index == 0 and reference is None:
            continue
        style_id = reference.get(w("val")) if reference is not None else "Normal"
        seen = set()
        while style_id in by_id and style_id not in seen:
            seen.add(style_id)
            style = by_id[style_id]
            candidates.append(style.find("w:rPr/w:rFonts", NS))
            base = style.find("w:basedOn", NS)
            style_id = base.get(w("val")) if base is not None else None
    candidates.append(styles.find("w:docDefaults/w:rPrDefault/w:rPr/w:rFonts", NS))
    for fonts in candidates:
        if fonts is None:
            continue
        theme_name = fonts.get(w(f"{slot}Theme"))
        if theme_name is not None:
            role = "major" if theme_name.startswith("major") else "minor"
            return theme.find(f"a:themeElements/a:fontScheme/a:{role}Font/a:latin", NS).get("typeface")
        literal = fonts.get(w(slot))
        if literal is not None:
            return literal
    raise AssertionError("No effective font found")


@pytest.mark.parametrize("body,heading", [("Aptos", "Aptos Display"), ("Arial", "Times New Roman")])
def test_saved_fonts_resolve_through_theme_without_direct_formatting(tmp_path: Path, body: str, heading: str) -> None:
    output = render_fonts(tmp_path, f"    body: {body}\n    headings: {heading}\n    monospace: Consolas\n")
    theme = xml(output, "word/theme/theme1.xml")
    styles = xml(output, "word/styles.xml")
    assert theme_names(theme) == (heading, body)
    content = xml(output, "word/document.xml")
    for paragraph in content.findall(".//w:p", NS):
        pstyle = paragraph.find("w:pPr/w:pStyle", NS)
        style_id = pstyle.get(w("val")) if pstyle is not None else "Normal"
        for run in paragraph.findall(".//w:r", NS):
            text = "".join(run.itertext())
            code = "inline code" in text or style_id == "CodeBlock"
            expected = "Consolas" if code else heading if style_id.startswith("Heading") else body
            for slot in ("ascii", "hAnsi"):
                assert effective_font(run, paragraph, styles, theme, slot) == expected, (text, slot)
            if not code:
                assert run.find("w:rPr/w:rFonts", NS) is None

    for style_id in [
        "Normal",
        "Quote",
        "ListBullet",
        "ListBullet2",
        "ListBullet3",
        "ListNumber",
        "ListNumber2",
        "ListNumber3",
    ]:
        fonts = styles.find(f"w:style[@w:styleId='{style_id}']/w:rPr/w:rFonts", NS)
        assert fonts.get(w("asciiTheme")) == "minorHAnsi"
        assert fonts.get(w("ascii")) is None
    for level in range(1, 7):
        for style_id in (f"Heading{level}", f"Heading{level}Char"):
            fonts = styles.find(f"w:style[@w:styleId='{style_id}']/w:rPr/w:rFonts", NS)
            assert fonts.get(w("asciiTheme")) == "majorHAnsi"
            assert fonts.get(w("hAnsiTheme")) == "majorHAnsi"
            assert fonts.get(w("ascii")) is None
            assert fonts.get(w("hAnsi")) is None

    defaults = styles.find("w:docDefaults/w:rPrDefault/w:rPr/w:rFonts", NS)
    assert defaults.get(w("asciiTheme")) == "minorHAnsi"
    assert defaults.get(w("hAnsiTheme")) == "minorHAnsi"
    assert defaults.get(w("ascii")) is None
    assert defaults.get(w("hAnsi")) is None

    # Changing only the theme must change effective output, including future headings.
    for role, font in (("major", "Georgia"), ("minor", "Verdana")):
        theme.find(f"a:themeElements/a:fontScheme/a:{role}Font/a:latin", NS).set("typeface", font)
    for paragraph in content.findall(".//w:p", NS):
        pstyle = paragraph.find("w:pPr/w:pStyle", NS)
        style_id = pstyle.get(w("val")) if pstyle is not None else "Normal"
        for run in paragraph.findall(".//w:r", NS):
            code = "inline code" in "".join(run.itertext()) or style_id == "CodeBlock"
            expected = "Consolas" if code else "Georgia" if style_id.startswith("Heading") else "Verdana"
            assert effective_font(run, paragraph, styles, theme, "ascii") == expected


@pytest.mark.parametrize(
    "override,changed,expected",
    [("body: Aptos", "minor", "Aptos"), ("headings: Aptos Display", "major", "Aptos Display")],
)
def test_partial_override_preserves_other_theme_fonts_and_styles(
    tmp_path: Path, override: str, changed: str, expected: str
) -> None:
    output = render_fonts(tmp_path, f"    {override}\n")
    with ZipFile(BytesIO(default_template_bytes())) as package:
        original_theme = ET.fromstring(package.read("word/theme/theme1.xml"))
        original_styles = ET.fromstring(package.read("word/styles.xml"))
    theme = xml(output, "word/theme/theme1.xml")
    styles = xml(output, "word/styles.xml")
    for role in ("major", "minor"):
        path = f"a:themeElements/a:fontScheme/a:{role}Font"
        before = original_theme.find(path, NS)
        after = theme.find(path, NS)
        if role == changed:
            assert after.find("a:latin", NS).get("typeface") == expected
            for tag in ("ea", "cs", "font"):
                assert [canonical(e) for e in before.findall(f"a:{tag}", NS)] == [
                    canonical(e) for e in after.findall(f"a:{tag}", NS)
                ]
        else:
            assert canonical(before) == canonical(after)
    unchanged_style = "Heading1" if changed == "minor" else "Normal"
    path = f"w:style[@w:styleId='{unchanged_style}']"
    assert canonical(original_styles.find(path, NS)) == canonical(styles.find(path, NS))
    for tag in ("clrScheme", "fmtScheme"):
        path = f"a:themeElements/a:{tag}"
        assert canonical(original_theme.find(path, NS)) == canonical(theme.find(path, NS))


def test_custom_style_mapping_preserves_formatting_and_template(tmp_path: Path) -> None:
    document = load_template(None)
    body = document.styles.add_style("Company Body", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = document.styles["Normal"]
    body.font.name = "Courier New"
    body.font.size = Pt(14)
    heading = document.styles.add_style("Company Heading", WD_STYLE_TYPE.PARAGRAPH)
    heading.base_style = document.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.bold = True
    heading.paragraph_format.keep_with_next = True
    template = tmp_path / "template.docx"
    document.save(template)
    before = template.read_bytes()
    output = render_fonts(
        tmp_path,
        "    body: Aptos\n    headings: Aptos Display\n",
        template=template,
        styles="  styles:\n    paragraph: Company Body\n    headings:\n      1: Company Heading\n",
    )
    saved = Document(output)
    assert saved.styles["Company Body"].font.name is None
    assert saved.styles["Company Body"].font.size == Pt(14)
    assert saved.styles["Company Heading"].font.name is None
    assert saved.styles["Company Heading"].font.bold is True
    assert saved.styles["Company Heading"].paragraph_format.keep_with_next is True
    styles = xml(output, "word/styles.xml")
    theme = xml(output, "word/theme/theme1.xml")
    content = xml(output, "word/document.xml")
    first = content.find("w:body/w:p", NS)
    assert effective_font(first.find("w:r", NS), first, styles, theme, "ascii") == "Aptos Display"
    assert template.read_bytes() == before


def test_template_without_theme_gets_valid_theme_relationship(tmp_path: Path) -> None:
    document = load_template(None)
    theme_rel = next(rel.rId for rel in document.part.rels.values() if rel.reltype == RELATIONSHIP_TYPE.THEME)
    document.part.drop_rel(theme_rel)
    template = tmp_path / "no-theme.docx"
    document.save(template)
    output = render_fonts(tmp_path, "    body: Aptos\n    headings: Aptos Display\n", template=template)
    saved = Document(output)
    part = saved.part.part_related_by(RELATIONSHIP_TYPE.THEME)
    assert theme_names(ET.fromstring(part.blob)) == ("Aptos Display", "Aptos")


@pytest.mark.parametrize("mapping", ["    headings:\n      1: Normal\n", "    code_block: Heading 1\n"])
def test_conflicting_font_style_roles_are_rejected(tmp_path: Path, mapping: str) -> None:
    with pytest.raises(TemplateError, match="distinct styles") as error:
        render_fonts(tmp_path, "    body: Aptos\n    headings: Aptos Display\n", styles="  styles:\n" + mapping)
    assert error.value.context.code == "template_font_style_conflict"
    assert not (tmp_path / "fonts.docx").exists()


def test_monospace_override_clears_conflicting_theme_reference(tmp_path: Path) -> None:
    output = render_fonts(tmp_path, "    monospace: Cascadia Mono\n", styles="  styles:\n    code_block: Heading 6\n")
    styles = xml(output, "word/styles.xml")
    fonts = styles.find("w:style[@w:styleId='Heading6']/w:rPr/w:rFonts", NS)
    assert fonts.get(w("ascii")) == "Cascadia Mono"
    assert fonts.get(w("asciiTheme")) is None
    assert fonts.get(w("hAnsiTheme")) is None
    with ZipFile(BytesIO(default_template_bytes())) as source, ZipFile(output) as rendered:
        assert rendered.read("word/theme/theme1.xml") == source.read("word/theme/theme1.xml")


@pytest.mark.parametrize(
    "theme_bytes", [b"<broken", b'<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>']
)
def test_malformed_theme_is_reported_without_replacing_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, theme_bytes: bytes
) -> None:
    original_blob = Part.blob.fget

    def blob(part: Part) -> bytes:
        if part.content_type == "application/vnd.openxmlformats-officedocument.theme+xml":
            return theme_bytes
        return original_blob(part)

    monkeypatch.setattr(Part, "blob", property(blob))
    output = tmp_path / "fonts.docx"
    output.write_bytes(b"existing output")
    with pytest.raises(TemplateError) as error:
        render_fonts(tmp_path, "    body: Aptos\n    headings: Aptos Display\n")
    assert error.value.context.code == "template_theme_invalid"
    assert output.read_bytes() == b"existing output"


def test_no_font_overrides_preserve_theme_and_heading_definitions(tmp_path: Path) -> None:
    output = render_fonts(tmp_path, "")
    with ZipFile(BytesIO(default_template_bytes())) as source, ZipFile(output) as saved:
        assert source.read("word/theme/theme1.xml") == saved.read("word/theme/theme1.xml")
        before = ET.fromstring(source.read("word/styles.xml"))
        after = ET.fromstring(saved.read("word/styles.xml"))
    for style_id in ("Normal", "Quote", *(f"Heading{level}" for level in range(1, 7))):
        path = f"w:style[@w:styleId='{style_id}']"
        assert canonical(before.find(path, NS)) == canonical(after.find(path, NS))


def test_new_paragraphs_and_linked_styles_inherit_theme_after_reopening(tmp_path: Path) -> None:
    output = render_fonts(tmp_path, "    body: Aptos\n    headings: Aptos Display\n")
    document = Document(output)
    document.add_paragraph("New body")
    for level in range(1, 7):
        document.add_heading(f"New heading {level}", level=level)
        document.add_paragraph().add_run(f"New linked heading {level}", style=f"Heading {level} Char")
    document.save(output)
    styles = xml(output, "word/styles.xml")
    theme = xml(output, "word/theme/theme1.xml")
    content = xml(output, "word/document.xml")
    checked = 0
    for paragraph in content.findall(".//w:p", NS):
        for run in paragraph.findall(".//w:r", NS):
            text = "".join(run.itertext())
            if not text.startswith("New "):
                continue
            assert run.find("w:rPr/w:rFonts", NS) is None
            expected = "Aptos" if text == "New body" else "Aptos Display"
            assert effective_font(run, paragraph, styles, theme, "ascii") == expected
            checked += 1
    assert checked == 13
