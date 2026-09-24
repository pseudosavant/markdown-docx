from __future__ import annotations

from hashlib import sha256
from pathlib import Path

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from markdown_docx.errors import ParseError
from markdown_docx.models import HeadingBlock, ParagraphBlock
from markdown_docx.parser import parse_document
from markdown_docx.renderer import render_docx
from markdown_docx.template import load_template


def test_heading_slugs_are_deterministic_and_collision_safe() -> None:
    model = parse_document(
        "# **Hello**, `World`!\n\n# Hello World\n\n# Hello World-1\n\n# Café 東京\n\n# !!!\n",
        input_path=None,
        source_name="input.md",
    )
    assert [block.anchor for block in model.blocks if isinstance(block, HeadingBlock)] == [
        "hello-world",
        "hello-world-1",
        "hello-world-1-1",
        "café-東京",
        "section",
    ]


def test_forward_and_percent_encoded_links_resolve() -> None:
    model = parse_document(
        "[Next](#caf%C3%A9-%E6%9D%B1%E4%BA%AC)\n\n# Café 東京\n", input_path=None, source_name="input.md"
    )
    first = model.blocks[0]
    assert isinstance(first, ParagraphBlock)
    assert first.fragments[0].anchor == "café-東京"


@pytest.mark.parametrize(
    "target,code",
    [("missing", "internal_link_unresolved"), ("%FF", "internal_link_invalid"), ("", "internal_link_unresolved")],
)
def test_unresolved_links_have_source_lines(target: str, code: str) -> None:
    with pytest.raises(ParseError) as error:
        parse_document(f"Intro\n\n[Target](#{target})\n", input_path=None, source_name="input.md")
    assert error.value.context.code == code
    assert error.value.context.line == 3
    assert error.value.context.input_path == "input.md"


def test_internal_links_render_in_all_supported_contexts(tmp_path: Path) -> None:
    source = """[Forward **link**](#details "Jump")

# Details

[Second](#details-1)

# Details

> [Quote](#details)

- [List](#details)

| [Header](#details) |
| --- |
| [Cell](#details-1) |
"""
    model = parse_document(source, input_path=None, source_name="input.md")
    output = tmp_path / "bookmarks.docx"
    render_docx(model, output, template_path=None, base_dir=tmp_path, allow_remote_images=False)
    document = Document(output)
    bookmarks = list(document.bookmarks)
    assert len(bookmarks) == 2
    assert [bookmark.paragraph.text for bookmark in bookmarks] == ["Details", "Details"]
    assert len({bookmark.name for bookmark in bookmarks}) == 2
    links = [link for paragraph in document.paragraphs for link in paragraph.hyperlinks]
    links += [link for row in document.tables[0].rows for cell in row.cells for link in cell.paragraphs[0].hyperlinks]
    assert len(links) == 6
    assert {link.fragment for link in links} == {bookmark.name for bookmark in bookmarks}
    assert all(link.address == "" for link in links)
    assert links[0].tooltip == "Jump"
    assert links[0].runs[-1].bold is True
    assert not any(rel.reltype == RT.HYPERLINK for rel in document.part.rels.values())


def test_template_bookmark_names_are_preserved(tmp_path: Path) -> None:
    template = load_template(None)
    paragraph = template.add_paragraph() if not template.paragraphs else template.paragraphs[0]
    reserved = "md_" + sha256(b"details").hexdigest()[:28]
    template.bookmarks.add(reserved, paragraph=paragraph)
    path = tmp_path / "template.docx"
    template.save(path)
    original = path.read_bytes()
    model = parse_document("# Details\n\n[Link](#details)\n", input_path=None, source_name="input.md")
    output = tmp_path / "out.docx"
    render_docx(model, output, template_path=path, base_dir=tmp_path, allow_remote_images=False)
    document = Document(output)
    assert {bookmark.name for bookmark in document.bookmarks} == {reserved, reserved + "_1"}
    assert document.paragraphs[-1].hyperlinks[0].fragment == reserved + "_1"
    assert path.read_bytes() == original
