from __future__ import annotations

from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import RGBColor
from docx.text.hyperlink import Hyperlink

from markdown_docx.cli import main
from markdown_docx.errors import UnsupportedFeatureError
from markdown_docx.models import ParagraphBlock
from markdown_docx.parser import parse_document
from markdown_docx.renderer import render_docx

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def render(tmp_path: Path, source: str, template: Path | None = None) -> Path:
    model = parse_document(source, input_path=tmp_path / "input.md", source_name="input.md")
    output = tmp_path / "links.docx"
    render_docx(model, output, template_path=template, base_dir=tmp_path, allow_remote_images=False)
    return output


def test_link_boundaries_and_formatting_survive_parsing() -> None:
    model = parse_document(
        'Before [**bold** *italic* `code`](https://example.com "Details") after.',
        input_path=None,
        source_name="input.md",
    )
    paragraph = model.blocks[0]
    assert isinstance(paragraph, ParagraphBlock)
    opening = paragraph.fragments[1]
    assert (opening.kind, opening.href, opening.title) == ("link_open", "https://example.com", "Details")
    assert paragraph.fragments[-2].kind == "link_close"
    assert paragraph.fragments[-1].text == " after."
    assert any(f.text == "bold" and f.bold for f in paragraph.fragments)
    assert any(f.text == "italic" and f.italic for f in paragraph.fragments)
    assert any(f.text == "code" and f.code for f in paragraph.fragments)


def test_native_hyperlink_relationship_formatting_and_tooltip(tmp_path: Path) -> None:
    url = "https://example.com/search?q=one&lang=en#results"
    output = render(tmp_path, f'Before [Café **bold** *italic* `code`]({url} "Tips & details") after.')
    document = Document(output)
    paragraph = document.paragraphs[0]
    assert paragraph.text == "Before Café bold italic code after."
    assert [run.text for run in paragraph.runs] == ["Before ", " after."]
    hyperlink = paragraph.hyperlinks[0]
    assert hyperlink.url == url
    assert hyperlink.text == "Café bold italic code"
    assert hyperlink.tooltip == "Tips & details"
    assert len(paragraph.hyperlinks) == 1
    assert any(run.text == "bold" and run.bold for run in hyperlink.runs)
    assert any(run.text == "italic" and run.italic for run in hyperlink.runs)
    assert any(run.text == "code" and run.font.name == "Consolas" for run in hyperlink.runs)
    assert all(run.style.name == "Hyperlink" for run in hyperlink.runs)
    assert document.styles["Hyperlink"].font.underline
    with ZipFile(output) as archive:
        body = ET.fromstring(archive.read("word/document.xml"))
        relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
    element = body.find(".//w:hyperlink", NS)
    assert element is not None
    assert element.get(f"{{{NS['w']}}}tooltip") == "Tips & details"
    relationship_id = element.get(f"{{{NS['r']}}}id")
    relationship = next(item for item in relationships if item.get("Id") == relationship_id)
    assert relationship.get("Target") == url
    assert relationship.get("TargetMode") == "External"
    assert relationship.get("Type") == NS["r"] + "/hyperlink"


def test_adjacent_links_keep_distinct_boundaries_and_surrounding_formatting(tmp_path: Path) -> None:
    output = render(tmp_path, "**before [one](https://example.com)[two](https://example.com) after**")
    paragraph = Document(output).paragraphs[0]
    assert paragraph.text == "before onetwo after"
    assert [link.text for link in paragraph.hyperlinks] == ["one", "two"]
    assert all(run.bold for run in paragraph.runs)
    assert all(run.bold for link in paragraph.hyperlinks for run in link.runs)
    assert [isinstance(item, Hyperlink) for item in paragraph.iter_inner_content()] == [False, True, True, False]


@pytest.mark.parametrize("prefix", ["", "# ", "> ", "- ", "1. "])
def test_links_in_text_block_contexts(tmp_path: Path, prefix: str) -> None:
    paragraph = Document(render(tmp_path, prefix + "[site](https://example.com)")).paragraphs[0]
    assert paragraph.hyperlinks[0].url == "https://example.com"
    assert paragraph.text == "site"


def test_links_in_table_headers_and_cells(tmp_path: Path) -> None:
    output = render(tmp_path, "| [Header](https://example.com) |\n| --- |\n| [Email](mailto:hello@example.com) |")
    table = Document(output).tables[0]
    assert table.cell(0, 0).paragraphs[0].hyperlinks[0].url == "https://example.com"
    assert table.cell(1, 0).paragraphs[0].hyperlinks[0].url == "mailto:hello@example.com"


@pytest.mark.parametrize(
    ("source", "url", "label"),
    [
        ("[site][ref]\n\n[ref]: https://example.com", "https://example.com", "site"),
        ("<https://example.com>", "https://example.com", "https://example.com"),
        ("<hello@example.com>", "mailto:hello@example.com", "hello@example.com"),
        ("[file](guide.pdf)", "guide.pdf", "file"),
        ("[space](<https://example.com/a b>)", "https://example.com/a%20b", "space"),
    ],
)
def test_standard_markdown_link_forms(tmp_path: Path, source: str, url: str, label: str) -> None:
    link = Document(render(tmp_path, source)).paragraphs[0].hyperlinks[0]
    assert (link.url, link.text) == (url, label)


def test_breaks_and_literal_code_in_link_labels(tmp_path: Path) -> None:
    paragraph = Document(render(tmp_path, "[first  \nsecond\nthird `[^1]`](https://example.com)")).paragraphs[0]
    assert paragraph.hyperlinks[0].text == "first\nsecond third [^1]"


def test_linked_inline_image_stays_clickable(tmp_path: Path, png_file: Path) -> None:
    output = render(tmp_path, f"[![image]({png_file.as_posix()})](https://example.com)")
    document = Document(output)
    assert document.paragraphs[0].hyperlinks[0].url == "https://example.com"
    with ZipFile(output) as archive:
        body = ET.fromstring(archive.read("word/document.xml"))
    assert body.find(".//w:hyperlink/w:r/w:drawing", NS) is not None


@pytest.mark.parametrize("destination", [""])
def test_unsupported_destinations_are_line_aware(destination: str) -> None:
    with pytest.raises(UnsupportedFeatureError) as excinfo:
        parse_document(f"Intro\n\n[link]({destination})", input_path=None, source_name="input.md")
    assert excinfo.value.context.code == "unsupported_feature"
    assert excinfo.value.context.line == 3
    assert excinfo.value.context.input_path == "input.md"


def test_existing_hyperlink_style_is_preserved(tmp_path: Path) -> None:
    document = Document()
    style = document.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    style.font.color.rgb = RGBColor.from_string("118833")
    style.font.underline = False
    document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    template = tmp_path / "template.docx"
    document.save(template)
    output = render(tmp_path, "[site](https://example.com)", template)
    reopened = Document(output)
    assert reopened.styles["Hyperlink"].font.color.rgb == RGBColor.from_string("118833")
    assert reopened.styles["Hyperlink"].font.underline is False
    assert reopened.paragraphs[0].hyperlinks[0].runs[0].style.name == "Hyperlink"


def test_cli_converts_markdown_links(tmp_path: Path) -> None:
    source = tmp_path / "input.md"
    source.write_text("[Link text](https://example.com)", encoding="utf-8")
    assert main([str(source)]) == 0
    link = Document(source.with_suffix(".docx")).paragraphs[0].hyperlinks[0]
    assert (link.text, link.url) == ("Link text", "https://example.com")
