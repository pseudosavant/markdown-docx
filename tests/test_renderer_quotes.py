from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.shared import Inches

from markdown_docx.models import (
    Block,
    CodeBlock,
    HeadingBlock,
    ImageBlock,
    ListParagraphBlock,
    ParagraphBlock,
    TableBlock,
)
from markdown_docx.parser import parse_document
from markdown_docx.renderer import render_docx


def _render(tmp_path: Path, source: str) -> tuple[list[Block], DocumentObject]:
    model = parse_document(source, input_path=tmp_path / "input.md", source_name="input.md")
    output = tmp_path / "quotes.docx"
    render_docx(model, output, template_path=None, base_dir=tmp_path, allow_remote_images=False)
    return model.blocks, Document(output)


def test_quote_contains_editable_heading_code_list_table_image_and_nested_quote(tmp_path: Path, png_file: Path) -> None:
    source = """> # Quoted heading
>
> Quoted paragraph with ![Inline](image.png).
>
> ```text
> code in quote
> ```
>
> 1. Numbered detail
> 2. Another detail
>
> | Name | Value |
> | --- | --- |
> | a | b |
>
> ![Chart](image.png)
>
> > Deep quote
"""
    blocks, document = _render(tmp_path, source)
    assert [type(block) for block in blocks] == [
        HeadingBlock,
        ParagraphBlock,
        CodeBlock,
        ListParagraphBlock,
        ListParagraphBlock,
        TableBlock,
        ImageBlock,
        ParagraphBlock,
    ]
    assert blocks[-1].quote_depth == 2
    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.paragraphs[0].paragraph_format.left_indent == Inches(0.5)
    assert document.paragraphs[1].style.name == "Quote"
    assert document.paragraphs[2].style.name == "Code Block"
    assert document.paragraphs[3].paragraph_format.left_indent == Inches(0.75)
    assert document.paragraphs[-1].paragraph_format.left_indent == Inches(1)
    assert document.tables[0].left_indent == Inches(0.5)
    assert document.tables[0].cell(1, 0).text == "a"
    assert [shape.description for shape in document.inline_shapes] == ["Inline", "Chart"]


def test_quote_inside_list_can_contain_code_table_image_and_nested_list(tmp_path: Path, png_file: Path) -> None:
    source = """1. Parent

   > Quoted detail
   >
   > ```text
   > code
   > ```
   >
   > - Child bullet
   >
   > | A | B |
   > | --- | --- |
   > | x | y |
   >
   > ![Nested picture](image.png)

   After quote

2. Next parent
"""
    _, document = _render(tmp_path, source)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Parent",
        "Quoted detail",
        "code",
        "Child bullet",
        "",
        "After quote",
        "Next parent",
    ]
    assert document.paragraphs[1].style.name == "Quote"
    assert document.paragraphs[1].paragraph_format.left_indent == Inches(0.75)
    assert document.paragraphs[2].paragraph_format.left_indent == Inches(0.75)
    assert document.paragraphs[3].style.name == "List Bullet 2"
    assert document.paragraphs[3].paragraph_format.left_indent == Inches(1)
    assert document.paragraphs[4].paragraph_format.left_indent == Inches(0.75)
    assert document.tables[0].left_indent == Inches(0.75)
    assert document.inline_shapes[0].description == "Nested picture"
    assert document.paragraphs[5]._p.pPr.numPr.numId.val == 0
    assert document.paragraphs[6]._p.pPr.numPr.numId.val == document.paragraphs[0]._p.pPr.numPr.numId.val
